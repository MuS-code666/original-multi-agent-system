"""slides.json から M&A 候補比較の pptx を生成する。

使い方:
    python build_pptx.py <slides.json> [出力先.pptx]
出力先を省略すると output/MA候補比較_{YYYYMMDD}.pptx に保存する。
"""
import json
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FONT = "Meiryo"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
TEXT = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CONFIDENCE_COLORS = {
    "確認済み": RGBColor(0x1E, 0x84, 0x49),
    "推定": RGBColor(0xB9, 0x77, 0x0E),
    "不明": RGBColor(0x80, 0x80, 0x80),
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
BODY_TOP = Inches(1.35)
BODY_W = SLIDE_W - MARGIN * 2
BODY_H = Inches(5.2)

REQUIRED_FIELDS = {
    "title": ["title"],
    "bullets": ["title", "bullets"],
    "table": ["title", "columns", "rows"],
    "two_column": ["title", "left", "right"],
    "company": ["title", "facts", "points"],
    "sources": ["title", "items"],
}


def validate(deck):
    errors = []
    if not isinstance(deck.get("slides"), list) or not deck["slides"]:
        return ["slides が空、または配列ではありません"]
    for i, slide in enumerate(deck["slides"], start=1):
        slide_type = slide.get("type")
        if slide_type not in REQUIRED_FIELDS:
            errors.append(f"スライド{i}: 不明な type '{slide_type}'（使えるのは {', '.join(REQUIRED_FIELDS)}）")
            continue
        for field in REQUIRED_FIELDS[slide_type]:
            if field not in slide:
                errors.append(f"スライド{i}（{slide_type}）: '{field}' がありません")
        if slide_type == "table" and "columns" in slide and "rows" in slide:
            for r, row in enumerate(slide["rows"], start=1):
                if len(row) != len(slide["columns"]):
                    errors.append(f"スライド{i}: {r}行目の列数が columns と一致しません")
        if slide_type == "company":
            for fact in slide.get("facts", []):
                if fact.get("confidence") not in CONFIDENCE_COLORS:
                    errors.append(f"スライド{i}: facts の confidence は 確認済み／推定／不明 のいずれかにしてください（{fact.get('label')}）")
    return errors


def style_run(run, size, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, left, top, width, height, lines, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines は文字列、または (文字列, 文字サイズ, 太字, 色) のタプルのリスト。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        text, line_size, line_bold, line_color = line if isinstance(line, tuple) else (line, size, bold, color)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(line_size * 0.4)
        run = p.add_run()
        style_run(run, line_size, line_bold, line_color)
        run.text = text
    return box


def bullet_lines(items, size):
    """items は文字列、または {"text": ..., "sub": [...]} のリスト。"""
    lines = []
    for item in items:
        if isinstance(item, dict):
            lines.append(("・" + item["text"], size, False, TEXT))
            lines.extend(("　　- " + sub, size - 2, False, GRAY) for sub in item.get("sub", []))
        else:
            lines.append(("・" + item, size, False, TEXT))
    return lines


def bullet_size(items):
    count = sum(1 + len(i.get("sub", [])) if isinstance(i, dict) else 1 for i in items)
    return 20 if count <= 4 else 18 if count <= 7 else 16 if count <= 10 else 14


def add_frame(prs, title, page, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.05))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    add_text(slide, MARGIN, 0, BODY_W, Inches(1.05), [title], 26, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.05), SLIDE_W, Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()
    add_text(slide, MARGIN, SLIDE_H - Inches(0.5), Inches(10), Inches(0.35), [footer], 10, GRAY)
    add_text(slide, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.5), Inches(1), Inches(0.35), [str(page)], 10, GRAY, align=PP_ALIGN.RIGHT)
    return slide


def add_note(slide, data):
    if data.get("note"):
        add_text(slide, MARGIN, SLIDE_H - Inches(0.95), BODY_W, Inches(0.4), ["※ " + data["note"]], 12, GRAY)
    if data.get("notes"):
        slide.notes_slide.notes_text_frame.text = data["notes"]


def fill_cell(cell, text, size, bold=False, color=TEXT, fill=None):
    cell.text = ""
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = cell.text_frame.paragraphs[0].add_run()
    style_run(run, size, bold, color)
    run.text = str(text)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def add_table(slide, left, top, width, columns, rows, col_widths=None, size=None, cell_colors=None):
    size = size or (14 if len(rows) <= 5 else 12 if len(rows) <= 8 else 10)
    row_h = Inches(0.45 if size >= 14 else 0.38)
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, row_h * (len(rows) + 1))
    table = shape.table
    ratios = col_widths or [1] * len(columns)
    for c, ratio in enumerate(ratios):
        table.columns[c].width = int(width * ratio / sum(ratios))
    for c, name in enumerate(columns):
        fill_cell(table.cell(0, c), name, size, True, WHITE, NAVY)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            color = (cell_colors or {}).get((r, c), TEXT)
            fill_cell(table.cell(r, c), value, size, False, color, LIGHT if r % 2 == 0 else WHITE)
    return shape


def render_title(prs, data, deck):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(3.55), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    add_text(slide, MARGIN, Inches(1.8), BODY_W, Inches(1.6), [data["title"]], 40, WHITE, True, anchor=MSO_ANCHOR.BOTTOM)
    lines = [(data["subtitle"], 20, False, WHITE)] if data.get("subtitle") else []
    lines.append((data.get("date", deck.get("date", "")), 16, False, RGBColor(0xC8, 0xD3, 0xE0)))
    add_text(slide, MARGIN, Inches(3.85), BODY_W, Inches(1.5), lines, 20)
    add_note(slide, {"notes": data.get("notes")})


def render_bullets(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    add_text(slide, MARGIN, BODY_TOP, BODY_W, BODY_H, bullet_lines(data["bullets"], bullet_size(data["bullets"])), 18)
    add_note(slide, data)


def render_table(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    add_table(slide, MARGIN, BODY_TOP + Inches(0.1), BODY_W, data["columns"], data["rows"], data.get("col_widths"))
    add_note(slide, data)


def render_two_column(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    col_w = (BODY_W - Inches(0.4)) / 2
    for i, column in enumerate([data["left"], data["right"]]):
        left = MARGIN + (col_w + Inches(0.4)) * i
        head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, BODY_TOP, col_w, Inches(0.5))
        head.fill.solid()
        head.fill.fore_color.rgb = ACCENT
        head.line.fill.background()
        add_text(slide, left + Inches(0.15), BODY_TOP, col_w, Inches(0.5), [column.get("heading", "")], 18, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
        items = column.get("bullets", [])
        add_text(slide, left, BODY_TOP + Inches(0.65), col_w, BODY_H - Inches(0.65), bullet_lines(items, min(bullet_size(items), 18)), 16)
    add_note(slide, data)


def render_company(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    top = BODY_TOP
    if data.get("summary"):
        add_text(slide, MARGIN, top, BODY_W, Inches(0.5), [data["summary"]], 16, NAVY, True)
        top += Inches(0.6)
    left_w = BODY_W * 0.55
    facts = data["facts"]
    rows = [[f["label"], f["value"], f["confidence"]] for f in facts]
    colors = {(r, 2): CONFIDENCE_COLORS[f["confidence"]] for r, f in enumerate(facts, start=1)}
    add_table(slide, MARGIN, top, left_w, ["項目", "内容", "確度"], rows, [1.1, 3, 0.9], 12 if len(rows) <= 8 else 10, colors)
    right_left = MARGIN + left_w + Inches(0.35)
    right_w = BODY_W - left_w - Inches(0.35)
    add_text(slide, right_left, top, right_w, Inches(0.4), [data.get("points_heading", "評価ポイント")], 16, ACCENT, True)
    points = data["points"]
    add_text(slide, right_left, top + Inches(0.5), right_w, BODY_H - Inches(1.1), bullet_lines(points, min(bullet_size(points), 16)), 14)
    add_note(slide, data)


def render_sources(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    items = data["items"]
    size = 12 if len(items) <= 12 else 10 if len(items) <= 20 else 9
    if len(items) <= 12:
        add_text(slide, MARGIN, BODY_TOP, BODY_W, BODY_H, items, size)
    else:
        half = (len(items) + 1) // 2
        col_w = (BODY_W - Inches(0.3)) / 2
        add_text(slide, MARGIN, BODY_TOP, col_w, BODY_H, items[:half], size)
        add_text(slide, MARGIN + col_w + Inches(0.3), BODY_TOP, col_w, BODY_H, items[half:], size)
    add_note(slide, data)


RENDERERS = {
    "bullets": render_bullets,
    "table": render_table,
    "two_column": render_two_column,
    "company": render_company,
    "sources": render_sources,
}


def build(deck, output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    footer = deck.get("footer") or f"公開情報に基づく初期評価（{deck.get('date', date.today().isoformat())}時点）"
    for page, data in enumerate(deck["slides"], start=1):
        if data["type"] == "title":
            render_title(prs, data, deck)
        else:
            RENDERERS[data["type"]](prs, data, page, footer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) not in (2, 3):
        print(__doc__)
        sys.exit(2)
    json_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2]) if len(sys.argv) == 3 else Path("output") / f"MA候補比較_{date.today():%Y%m%d}.pptx"
    deck = json.loads(json_path.read_text(encoding="utf-8"))
    errors = validate(deck)
    if errors:
        print("slides.json にエラーがあります:")
        for error in errors:
            print("  - " + error)
        sys.exit(1)
    build(deck, output_path)
    print(f"生成しました: {output_path}（{len(deck['slides'])}枚）")


if __name__ == "__main__":
    main()
