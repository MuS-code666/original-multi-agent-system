"""スライド原稿（JSON）から、営業訪問前レポートの pptx を生成する。

使い方:
    python build_pptx.py <原稿.json> [出力先.pptx]
出力先を省略すると output/訪問前レポート_{原稿のファイル名}_{YYYYMMDD}.pptx に保存する。
"""
import json
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FONT = "Meiryo"
GREEN = RGBColor(0x1E, 0x5E, 0x3B)
ACCENT = RGBColor(0x3A, 0x9D, 0x5D)
TEXT = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xF1, 0xF6, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CONFIDENCE_COLORS = {
    "公表情報あり": RGBColor(0x1E, 0x6F, 0xB8),
    "業種から推定": RGBColor(0xB9, 0x77, 0x0E),
    "要ヒアリング": RGBColor(0xC0, 0x39, 0x2B),
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
BODY_TOP = Inches(1.35)
BODY_W = SLIDE_W - MARGIN * 2
BODY_H = Inches(5.2)

REQUIRED_FIELDS = {
    "title": ["title"],
    "cards": ["title", "cards"],
    "bullets": ["title", "bullets"],
    "table": ["title", "columns", "rows"],
    "two_column": ["title", "left", "right"],
    "company": ["title", "facts", "points"],
    "sources": ["title", "items"],
}


def validate(deck):
    errors = []
    if not isinstance(deck.get("slides"), list) or not deck["slides"]:
        return ["slides が空、または配列ではありません"]
    for i, slide in enumerate(deck["slides"], start=1):
        slide_type = slide.get("type")
        if slide_type not in REQUIRED_FIELDS:
            errors.append(f"スライド{i}: 不明な type '{slide_type}'（使えるのは {', '.join(REQUIRED_FIELDS)}）")
            continue
        for field in REQUIRED_FIELDS[slide_type]:
            if field not in slide:
                errors.append(f"スライド{i}（{slide_type}）: '{field}' がありません")
        if slide_type == "table" and "columns" in slide and "rows" in slide:
            for r, row in enumerate(slide["rows"], start=1):
                if len(row) != len(slide["columns"]):
                    errors.append(f"スライド{i}: {r}行目の列数が columns と一致しません")
            col = slide.get("confidence_col")
            if col is not None:
                for r, row in enumerate(slide["rows"], start=1):
                    if col < len(row) and row[col] not in CONFIDENCE_COLORS:
                        errors.append(f"スライド{i}: {r}行目の確度は {'／'.join(CONFIDENCE_COLORS)} のいずれかにしてください（'{row[col]}'）")
        if slide_type == "cards" and not 1 <= len(slide.get("cards", [])) <= 4:
            errors.append(f"スライド{i}: cards は1〜4枚にしてください")
    return errors


def style_run(run, size, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, left, top, width, height, lines, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines は文字列、または (文字列, 文字サイズ, 太字, 色) のタプルのリスト。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        text, line_size, line_bold, line_color = line if isinstance(line, tuple) else (line, size, bold, color)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(line_size * 0.4)
        run = p.add_run()
        style_run(run, line_size, line_bold, line_color)
        run.text = text
    return box


def add_rect(slide, left, top, width, height, color):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    return rect


def bullet_lines(items, size):
    """items は文字列、または {"text": ..., "sub": [...]} のリスト。"""
    lines = []
    for item in items:
        if isinstance(item, dict):
            lines.append(("・" + item["text"], size, False, TEXT))
            lines.extend(("　　- " + sub, size - 2, False, GRAY) for sub in item.get("sub", []))
        else:
            lines.append(("・" + item, size, False, TEXT))
    return lines


def bullet_size(items):
    count = sum(1 + len(i.get("sub", [])) if isinstance(i, dict) else 1 for i in items)
    return 20 if count <= 4 else 18 if count <= 7 else 16 if count <= 10 else 14


def add_frame(prs, title, page, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.05), GREEN)
    add_text(slide, MARGIN, 0, BODY_W, Inches(1.05), [title], 26, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, 0, Inches(1.05), SLIDE_W, Inches(0.06), ACCENT)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.5), Inches(10), Inches(0.35), [footer], 10, GRAY)
    add_text(slide, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.5), Inches(1), Inches(0.35), [str(page)], 10, GRAY, align=PP_ALIGN.RIGHT)
    return slide


def add_note(slide, data):
    if data.get("note"):
        add_text(slide, MARGIN, SLIDE_H - Inches(0.95), BODY_W, Inches(0.4), ["※ " + data["note"]], 12, GRAY)
    if data.get("notes"):
        slide.notes_slide.notes_text_frame.text = data["notes"]


def fill_cell(cell, text, size, bold=False, color=TEXT, fill=None):
    cell.text = ""
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = cell.text_frame.paragraphs[0].add_run()
    style_run(run, size, bold, color)
    run.text = str(text)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def add_table(slide, left, top, width, columns, rows, col_widths=None, size=None, confidence_col=None):
    size = size or (14 if len(rows) <= 5 else 12 if len(rows) <= 8 else 10)
    row_h = Inches(0.45 if size >= 14 else 0.38)
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, row_h * (len(rows) + 1))
    table = shape.table
    ratios = col_widths or [1] * len(columns)
    for c, ratio in enumerate(ratios):
        table.columns[c].width = int(width * ratio / sum(ratios))
    for c, name in enumerate(columns):
        fill_cell(table.cell(0, c), name, size, True, WHITE, GREEN)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            is_confidence = c == confidence_col
            color = CONFIDENCE_COLORS.get(value, TEXT) if is_confidence else TEXT
            fill_cell(table.cell(r, c), value, size, is_confidence, color, LIGHT if r % 2 == 0 else WHITE)
    return shape


def render_title(prs, data, deck):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, GREEN)
    add_rect(slide, MARGIN, Inches(3.55), Inches(1.2), Inches(0.08), ACCENT)
    add_text(slide, MARGIN, Inches(1.8), BODY_W, Inches(1.6), [data["title"]], 40, WHITE, True, anchor=MSO_ANCHOR.BOTTOM)
    lines = [(data["subtitle"], 20, False, WHITE)] if data.get("subtitle") else []
    lines.append((data.get("date", deck.get("date", "")), 16, False, RGBColor(0xCD, 0xE3, 0xD4)))
    add_text(slide, MARGIN, Inches(3.85), BODY_W, Inches(1.5), lines, 20)
    add_note(slide, {"notes": data.get("notes")})


def render_cards(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    cards = data["cards"]
    gap = Inches(0.3)
    card_w = (BODY_W - gap * (len(cards) - 1)) / len(cards)
    card_h = Inches(3.6)
    for i, card in enumerate(cards):
        left = MARGIN + (card_w + gap) * i
        add_rect(slide, left, BODY_TOP + Inches(0.1), card_w, card_h, LIGHT)
        add_rect(slide, left, BODY_TOP + Inches(0.1), card_w, Inches(0.08), ACCENT)
        add_text(slide, left + Inches(0.2), BODY_TOP + Inches(0.35), card_w - Inches(0.4), Inches(0.6), [card["heading"]], 18, GREEN, True)
        add_text(slide, left + Inches(0.2), BODY_TOP + Inches(1.0), card_w - Inches(0.4), card_h - Inches(1.1), [card["text"]], 15)
    add_note(slide, data)


def render_bullets(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    add_text(slide, MARGIN, BODY_TOP, BODY_W, BODY_H, bullet_lines(data["bullets"], bullet_size(data["bullets"])), 18)
    add_note(slide, data)


def render_table(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    add_table(slide, MARGIN, BODY_TOP + Inches(0.1), BODY_W, data["columns"], data["rows"], data.get("col_widths"), confidence_col=data.get("confidence_col"))
    add_note(slide, data)


def render_two_column(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    col_w = (BODY_W - Inches(0.4)) / 2
    for i, column in enumerate([data["left"], data["right"]]):
        left = MARGIN + (col_w + Inches(0.4)) * i
        add_rect(slide, left, BODY_TOP, col_w, Inches(0.5), ACCENT)
        add_text(slide, left + Inches(0.15), BODY_TOP, col_w, Inches(0.5), [column.get("heading", "")], 18, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
        items = column.get("bullets", [])
        add_text(slide, left, BODY_TOP + Inches(0.65), col_w, BODY_H - Inches(0.65), bullet_lines(items, min(bullet_size(items), 18)), 16)
    add_note(slide, data)


def render_company(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    top = BODY_TOP
    if data.get("summary"):
        add_text(slide, MARGIN, top, BODY_W, Inches(0.5), [data["summary"]], 16, GREEN, True)
        top += Inches(0.6)
    left_w = BODY_W * 0.55
    rows = [[f["label"], f["value"]] for f in data["facts"]]
    add_table(slide, MARGIN, top, left_w, ["項目", "内容"], rows, [1, 3], 12 if len(rows) <= 8 else 10)
    right_left = MARGIN + left_w + Inches(0.35)
    right_w = BODY_W - left_w - Inches(0.35)
    add_text(slide, right_left, top, right_w, Inches(0.4), [data.get("points_heading", "営業の着眼点")], 16, ACCENT, True)
    points = data["points"]
    add_text(slide, right_left, top + Inches(0.5), right_w, BODY_H - Inches(1.1), bullet_lines(points, min(bullet_size(points), 16)), 14)
    add_note(slide, data)


def render_sources(prs, data, page, footer):
    slide = add_frame(prs, data["title"], page, footer)
    items = data["items"]
    size = 12 if len(items) <= 12 else 10 if len(items) <= 20 else 9
    if len(items) <= 12:
        add_text(slide, MARGIN, BODY_TOP, BODY_W, BODY_H, items, size)
    else:
        half = (len(items) + 1) // 2
        col_w = (BODY_W - Inches(0.3)) / 2
        add_text(slide, MARGIN, BODY_TOP, col_w, BODY_H, items[:half], size)
        add_text(slide, MARGIN + col_w + Inches(0.3), BODY_TOP, col_w, BODY_H, items[half:], size)
    add_note(slide, data)


RENDERERS = {
    "cards": render_cards,
    "bullets": render_bullets,
    "table": render_table,
    "two_column": render_two_column,
    "company": render_company,
    "sources": render_sources,
}


def build(deck, output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    footer = deck.get("footer") or f"社内資料｜営業訪問前調査（{deck.get('date', date.today().isoformat())}時点）"
    for page, data in enumerate(deck["slides"], start=1):
        if data["type"] == "title":
            render_title(prs, data, deck)
        else:
            RENDERERS[data["type"]](prs, data, page, footer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) not in (2, 3):
        print(__doc__)
        sys.exit(2)
    json_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2]) if len(sys.argv) == 3 else Path("output") / f"訪問前レポート_{json_path.stem}_{date.today():%Y%m%d}.pptx"
    deck = json.loads(json_path.read_text(encoding="utf-8"))
    errors = validate(deck)
    if errors:
        print("スライド原稿にエラーがあります:")
        for error in errors:
            print("  - " + error)
        sys.exit(1)
    build(deck, output_path)
    print(f"生成しました: {output_path}（{len(deck['slides'])}枚）")


if __name__ == "__main__":
    main()
